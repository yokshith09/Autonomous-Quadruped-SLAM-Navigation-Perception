"""Export a report built as simple HTML blocks (headings, paragraphs, lists, tables, figures) to a native Word document.

Used by project_report.py so the Word and PDF versions come from the same content. Headings become Word heading styles,
tables become Word tables (header row repeated, rows kept whole), figures are embedded, and the footer carries page numbers.
`export` raises if the number of headings, tables or figures in the .docx differs from the HTML.
"""
import html as _html
import os
import re
from html.parser import HTMLParser

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

COLORS = {"good": "1B7F3B", "improve": "A35A00", "poor": "B3261E", "info": "555555", "note": "555555", "cap": "333333"}
BODY_PT, TABLE_PT, NOTE_PT = 10, 8, 8.5


class Inline(HTMLParser):
    """Adds the inline content of a block (bold, italic, line breaks, coloured spans) to a Word paragraph."""

    def __init__(self, par, size, base_color=None):
        super().__init__(convert_charrefs=True)
        self.par, self.size, self.base_color = par, size, base_color
        self.bold = self.ital = 0
        self.colors = []
        self.started = False

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in ("b", "strong"):
            self.bold += 1
        elif tag in ("i", "em"):
            self.ital += 1
        elif tag == "br":
            self.par.add_run().add_break()
        elif tag == "span":
            c = COLORS.get((a.get("class") or "").strip())
            self.colors.append(c)
            if (a.get("class") or "").strip() in ("good", "improve", "poor"):
                self.bold += 1
                self.colors[-1] = (c, True)

    def handle_endtag(self, tag):
        if tag in ("b", "strong"):
            self.bold = max(0, self.bold - 1)
        elif tag in ("i", "em"):
            self.ital = max(0, self.ital - 1)
        elif tag == "span" and self.colors:
            c = self.colors.pop()
            if isinstance(c, tuple):
                self.bold = max(0, self.bold - 1)

    def handle_data(self, data):
        text = re.sub(r"\s+", " ", data)
        if not self.started:
            text = text.lstrip()
        if not text:
            return
        self.started = True
        run = self.par.add_run(text)
        run.font.size = Pt(self.size)
        run.bold = bool(self.bold) or None
        run.italic = bool(self.ital) or None
        c = next((x[0] if isinstance(x, tuple) else x for x in reversed(self.colors) if x), None) or self.base_color
        if c:
            run.font.color.rgb = RGBColor.from_string(c)


def add_inline(par, markup, size, base_color=None):
    p = Inline(par, size, base_color)
    p.feed(markup)
    p.close()


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def row_flags(row, header):
    trPr = row._tr.get_or_add_trPr()
    keep = OxmlElement("w:cantSplit")
    trPr.append(keep)
    if header:
        h = OxmlElement("w:tblHeader")
        trPr.append(h)


def add_table(doc, block, text_width_cm):
    rows = re.findall(r"<tr>(.*?)</tr>", block, re.S)
    cells = [re.findall(r"<td([^>]*)>(.*?)</td>", r, re.S) for r in rows]
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    head = cells[0]
    w = [int(m.group(1)) if (m := re.search(r"width=(\d+)", a)) else None for a, _ in head]
    if None in w or len(w) != ncol:
        w = [1] * ncol
    total = float(sum(w))
    widths = [Cm(text_width_cm * x / total) for x in w]
    for ri, row in enumerate(cells):
        is_head = "hdr" in row[0][0] if row else False
        row_flags(t.rows[ri], is_head)
        for ci in range(ncol):
            cell = t.cell(ri, ci)
            cell.width = widths[ci]
            attrs, content = row[ci] if ci < len(row) else ("", "")
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            if is_head:
                shade(cell, "E8EEF9")
                content = f"<b>{content}</b>"
            add_inline(par, content, TABLE_PT)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    return t


def add_page_number(par):
    run = par.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(8)


def export(html_text, img_dir, out, footer, title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    text_width_cm = 21.0 - 4.0
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        s = doc.styles[name]
        s.font.name, s.font.size = "Calibri", Pt(size)
        s.font.color.rgb = RGBColor(0x0B, 0x3D, 0x91)
    doc.core_properties.title = title
    doc.core_properties.author = ""
    doc.core_properties.comments = "Generated from the repository by evaluation/project_report.py"

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{footer}   |   page ")
    r.font.size = Pt(8)
    add_page_number(fp)

    counts = {"headings": 0, "tables": 0, "figures": 0}
    first_h1 = True
    for m in re.finditer(r"<(h[1-3]|p|table|ul)\b([^>]*)>(.*?)</\1>", html_text, re.S):
        tag, attrs, body = m.group(1), m.group(2), m.group(3)
        cls = (re.search(r"class=['\"]?(\w+)", attrs) or [None, ""])[1]
        if tag.startswith("h"):
            counts["headings"] += 1
            text = _html.unescape(re.sub(r"<[^>]+>", "", body)).strip()
            if tag == "h1" and first_h1:
                doc.add_heading(text, 0)
                first_h1 = False
            else:
                doc.add_heading(text, {"h1": 1, "h2": 1, "h3": 2}[tag])
        elif tag == "p":
            img = re.search(r"<img[^>]*src=['\"]([^'\"]+)['\"][^>]*?(?:width=['\"]?(\d+))?[^>]*/?>", body)
            if img:
                counts["figures"] += 1
                wpt = int(re.search(r"width=['\"]?(\d+)", img.group(0)).group(1)) if re.search(r"width=['\"]?(\d+)", img.group(0)) else 450
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.keep_with_next = True
                par.add_run().add_picture(os.path.join(img_dir, img.group(1)), width=Inches(min(wpt / 72.0, 6.6)))
                continue
            par = doc.add_paragraph(style="Caption" if cls == "cap" else None)
            if cls == "cap":
                par.paragraph_format.keep_with_next = True
            size = NOTE_PT if cls in ("cap", "note") else BODY_PT
            add_inline(par, body, size, COLORS.get(cls) if cls in ("cap", "note") else None)
        elif tag == "ul":
            for li in re.findall(r"<li>(.*?)</li>", body, re.S):
                add_inline(doc.add_paragraph(style="List Bullet"), li, BODY_PT)
        elif tag == "table":
            counts["tables"] += 1
            add_table(doc, m.group(0), text_width_cm)
    doc.save(out)

    # completeness check, as for the PDF
    got = Document(out)
    n_head = sum(1 for p in got.paragraphs if p.style.name.startswith(("Heading", "Title")))
    n_fig = len(got.inline_shapes)
    expect = {"headings": counts["headings"], "tables": counts["tables"], "figures": counts["figures"]}
    actual = {"headings": n_head, "tables": len(got.tables), "figures": n_fig}
    if expect != actual:
        raise RuntimeError(f"{out}: Word export incomplete, expected {expect}, found {actual}")
    print(f"  [SAVED] {out} ({actual['headings']} headings, {actual['tables']} tables, {actual['figures']} figures)")
    return out
